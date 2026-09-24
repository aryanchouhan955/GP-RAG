import os
import pytest
from docx import Document
from app.ingestion.loaders.docx_loader import load_docx

@pytest.fixture
def sample_docx(tmp_path):
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    
    doc.add_paragraph("This is paragraph 1.")
    doc.add_paragraph("   ")  # Empty after clean
    doc.add_paragraph("This is paragraph 3.")
    
    doc.save(str(docx_path))
    return str(docx_path)

def test_load_docx_valid(sample_docx):
    pages = load_docx(sample_docx)
    
    assert len(pages) == 2
    
    assert pages[0].page_number == 1
    assert "paragraph 1" in pages[0].text.lower()
    assert pages[0].metadata["file_type"] == "docx"
    
    assert pages[1].page_number == 3
    assert "paragraph 3" in pages[1].text.lower()

def test_load_docx_invalid(tmp_path):
    invalid_path = tmp_path / "nonexistent.docx"
    pages = load_docx(str(invalid_path))
    assert pages == []
